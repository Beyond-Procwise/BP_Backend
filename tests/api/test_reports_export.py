"""POST /reports/export: the builder's page becomes a real PDF / workbook.

The page comes from a caller, so the guard that matters most is that rendering it
cannot make this server fetch anything.
"""
import io

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from api.routers import reports as rr
from src.services import report_export

PAGE = "<html><body><h1>Q1 Review</h1><p>Demo</p></body></html>"


class _P:
    subject = "sub-real-caller"


@pytest.fixture
def client(monkeypatch):
    calls = {"gate": [], "stored": [], "recorded": []}
    monkeypatch.setattr(rr, "gate", lambda action, principal, **kw: calls["gate"].append(action))
    monkeypatch.setattr(report_export, "store",
                        lambda c, k, t: calls["stored"].append((k, t)) or True)
    monkeypatch.setattr(report_export, "record_url",
                        lambda rid, k: calls["recorded"].append((rid, k)) or True)
    app = FastAPI()
    app.include_router(rr.router)
    app.dependency_overrides[rr.require_user] = lambda: _P()
    c = TestClient(app)
    c.calls = calls
    return c


def test_a_page_becomes_a_real_pdf_and_is_kept(client):
    r = client.post("/reports/export", json={"html": PAGE, "name": "Q1 Review", "report_id": "42"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert 'attachment; filename="Q1 Review.pdf"' in r.headers["content-disposition"]
    assert r.headers["x-report-stored"] == "true" and r.headers["x-report-url-recorded"] == "true"
    assert client.calls["gate"] == ["report.read"]  # your own report: not a share
    assert client.calls["recorded"][0][0] == "42"


def test_tables_become_a_workbook_one_sheet_each(client):
    from openpyxl import load_workbook
    r = client.post("/reports/export", json={"format": "xlsx", "name": "t", "tables": [
        {"title": "Spend", "columns": ["m", "v"], "rows": [["Jan", 1], ["Feb", "=1+1"]]}]})
    ws = load_workbook(io.BytesIO(r.content))["Spend"]
    assert [c.value for c in ws[1]] == ["m", "v"] and ws["B3"].value == "'=1+1"  # never a live formula


@pytest.mark.parametrize("url", ["file:///etc/passwd", "http://169.254.169.254/latest/meta-data/",
                                 "http://localhost:8000/admin"])
def test_the_page_cannot_make_the_server_fetch(client, url):
    r = client.post("/reports/export", json={"html": f'<html><body><img src="{url}"></body></html>'})
    assert r.status_code == 422 and "outside resource" in r.json()["detail"]
    r = client.post("/reports/export", json={"html": PAGE, "css": f'body{{background:url("{url}")}}'})
    assert r.status_code == 422


def test_a_data_uri_image_still_renders(client):
    png = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGP4z8DwHwAFAAH/"
           "iZk9HQAAAABJRU5ErkJggg==")
    r = client.post("/reports/export", json={"html": f'<html><body><img src="{png}"></body></html>'})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_an_empty_page_is_refused_not_rendered_blank(client):
    assert client.post("/reports/export", json={"html": "  "}).status_code == 422
    assert client.post("/reports/export", json={"format": "xlsx", "tables": []}).status_code == 422


def test_a_refusal_by_the_gate_produces_no_file(monkeypatch, client):
    def deny(*a, **k):
        raise HTTPException(status_code=403, detail="refused")
    monkeypatch.setattr(rr, "gate", deny)
    assert client.post("/reports/export", json={"html": PAGE}).status_code == 403
    assert client.calls["stored"] == []


def test_a_storage_failure_still_downloads_and_says_so(monkeypatch, client):
    monkeypatch.setattr(report_export, "store", lambda *a: False)
    r = client.post("/reports/export", json={"html": PAGE, "report_id": "42"})
    assert r.status_code == 200 and r.headers["x-report-stored"] == "false"
    assert r.headers["x-report-url-recorded"] == "false"
