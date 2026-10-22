"""The export builders: real files from plain content, nothing invented, limits refused by name."""
import io

import pytest

from src.services import export_builders as EB

SECTIONS = [{"heading": "Position", "paragraphs": ["Spend with Acme was 12,400 GBP."], "bullets": ["One finding open."],
             "table": {"columns": ["Supplier", "Spend"], "rows": [["Acme", 12400], ["Beta", "3,200.50"]]}},
            {"heading": "Next steps", "paragraphs": ["Ask finance to confirm the budget."]}]


def test_docx_holds_the_title_sections_bullets_table_and_note():
    from docx import Document
    doc = Document(io.BytesIO(EB.docx_bytes("Budget note", SECTIONS, notes="Figures from the spend cube.", prepared_for="carol")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Budget note" in text and "Prepared for carol." in text and "One finding open." in text and "Figures from the spend cube." in text
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Supplier", "Spend"] and doc.tables[0].rows[2].cells[1].text == "3,200.50"


def test_xlsx_keeps_numbers_as_numbers_and_names_sheets_safely():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(EB.xlsx_bytes("Spend", [
        {"name": "By supplier", "columns": ["Supplier", "Spend", "Invoices"], "rows": [["Acme", "12,400.50", 3], ["Beta", 3200, None]]},
        {"name": "Bad/Name?", "columns": ["x"], "rows": [[1]]}, {"name": "Bad/Name?", "columns": ["x"], "rows": [[2]]}])))
    ws = wb["By supplier"]
    assert ws["B2"].value == 12400.5 and ws["C2"].value == 3 and ws["B3"].value == 3200 and ws["C3"].value is None
    assert ws.freeze_panes == "A2" and wb.sheetnames[1:] == ["Bad Name ", "Bad Name  2"]


def test_pptx_has_a_title_slide_then_one_slide_per_entry_with_tables_and_notes():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(EB.pptx_bytes("Award memo", [
        {"title": "Recommendation", "bullets": ["Award to B", "A is the alternative"], "notes": "close call"},
        {"title": "Scores", "table": {"columns": ["Bidder", "Score"], "rows": [["A", 0.81], ["B", 0.84]]}}], subtitle="Packaging tender")))
    assert len(prs.slides) == 3
    assert prs.slides[1].notes_slide.notes_text_frame.text == "close call"
    table = next(sh for sh in prs.slides[2].shapes if sh.has_table).table
    assert table.cell(0, 0).text == "Bidder" and table.cell(2, 1).text == "0.84"


@pytest.mark.parametrize("kind", ["bar", "line", "donut"])
def test_each_chart_type_renders_every_label_and_value_as_svg(kind):
    data = [{"label": "Acme", "value": 12400}, {"label": "Beta", "value": 3200.5}, {"label": "Gamma", "value": 800}]
    svg = EB.chart_svg(kind, "Spend by supplier", data, unit="GBP").decode()
    assert svg.startswith("<svg") and "Spend by supplier" in svg
    for label in ("Acme", "Beta", "Gamma"):
        assert label in svg
    assert "12,400 GBP" in svg and "3,200.50 GBP" in svg


def test_labels_are_escaped_so_a_chart_cannot_carry_markup():
    svg = EB.chart_svg("bar", "<script>", [{"label": "a<b", "value": 1}]).decode()
    assert "<script>" not in svg and "&lt;script&gt;" in svg and "a&lt;b" in svg


@pytest.mark.parametrize("call, needle", [
    (lambda: EB.docx_bytes("t", []), "at least one section"),
    (lambda: EB.docx_bytes("t", [{"table": {"rows": [[1]]}}]), "needs columns"),
    (lambda: EB.docx_bytes("t", [{"table": {"columns": ["a"], "rows": [[1]] * (EB.MAX_TABLE_ROWS + 1)}}]), "the cap is"),
    (lambda: EB.xlsx_bytes("t", []), "at least one sheet"),
    (lambda: EB.pptx_bytes("t", [{"title": "s", "bullets": ["b"] * (EB.MAX_BULLETS + 1)}]), "bullets; the cap"),
    (lambda: EB.chart_svg("pie", "t", [{"label": "a", "value": 1}]), "chart_type must be one of"),
    (lambda: EB.chart_svg("donut", "t", [{"label": "a", "value": -1}]), "cannot show a negative"),
    (lambda: EB.chart_svg("bar", "t", [{"label": "a", "value": "many"}]), "is not a number"),
    (lambda: EB.chart_svg("bar", "t", [{"label": "a", "value": 1}] * (EB.MAX_CHART_POINTS + 1)), "points; the cap"),
])
def test_content_outside_the_limits_is_refused_with_the_reason(call, needle):
    with pytest.raises(EB.ExportRefused) as exc:
        call()
    assert needle in str(exc.value)


def test_a_short_row_is_padded_and_a_long_row_is_cut_to_the_columns():
    from docx import Document
    doc = Document(io.BytesIO(EB.docx_bytes("t", [{"table": {"columns": ["a", "b"], "rows": [[1], [1, 2, 3]]}}])))
    rows = doc.tables[0].rows
    assert [c.text for c in rows[1].cells] == ["1", ""] and [c.text for c in rows[2].cells] == ["1", "2"]


def test_kpi_card_shows_value_movement_target_source_and_the_stated_status_colour():
    svg = EB.kpi_card_svg("On-time delivery", 92.4, unit="%", movement_pct=-1.5, target="95%", status="amber", source="scorecard, FY to date").decode()
    assert svg.startswith("<svg") and "On-time delivery" in svg and "92.40 %" in svg and "-1.5% versus prior period" in svg
    assert "Target 95%" in svg and "Source: scorecard, FY to date" in svg and EB.STATUS_COLOURS["amber"] in svg


@pytest.mark.parametrize("kwargs, needle", [
    ({"value": None}, "needs a value"),
    ({"value": 1, "status": "blue"}, "status must be one of"),
    ({"value": 1, "movement_pct": "up"}, "movement_pct must be a number"),
])
def test_kpi_card_refuses_bad_input_by_name(kwargs, needle):
    with pytest.raises(EB.ExportRefused) as exc:
        EB.kpi_card_svg("t", **kwargs)
    assert needle in str(exc.value)
