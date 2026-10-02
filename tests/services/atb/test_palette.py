from src.services.atb.pptx_import.evidence import Evidence
from src.services.atb.pptx_import.palette import census, colours
from src.services.atb.pptx_import.read import read_deck


def _deck(slides=(), theme=None):
    from src.services.atb.pptx_import.read import Deck
    return Deck(width_in=13.333, height_in=7.5, slides=tuple(slides),
                theme=theme or {'colours': {}, 'fonts': {}}, chart_series_colours=(), run_langs={})


def _uncoloured_runs(n=12):
    from src.services.atb.pptx_import.read import Box, Run, Shape
    runs = (Run(text='words', size_pt=12, font='Calibri', colour=None, bold=False, lang='en-GB'),)
    return tuple((Shape(kind='text', box=Box(0.5, 1.5, 3, 0.4), runs=runs, fill=None, line=None,
                        table=None, slide=i, name='t'),) for i in range(1, n + 1))


def test_the_theme_dark_colour_is_the_ink_of_last_resort():
    # Reached only when no run in the deck states a colour. The stock PowerPoint template's dk1 is
    # a sysClr, so no built fixture can exercise this: the Deck is constructed directly.
    ev = Evidence()
    out = colours(_deck(_uncoloured_runs(), theme={'colours': {'dk1': '#123456'}, 'fonts': {}}), ev)
    assert out['ink'] == '#123456'
    record = ev.as_dict()['values']['colours.ink']
    assert record['assumed'] is True, 'a theme colour is not a measurement of the deck'
    assert "theme's dk1" in record['why']
    # A deck that states no colour at all assumes its ink, its muted, its panel AND its accent —
    # and every one of them is a problem the review screen must show, not a note in the evidence.
    assert [a['path'] for a in ev.assumptions] == [
        'colours.ink', 'colours.muted', 'colours.panel', 'colours.accent']


def test_without_a_theme_dark_colour_the_ink_is_black_and_says_so():
    ev = Evidence()
    out = colours(_deck(_uncoloured_runs()), ev)
    assert out['ink'] == '#000000'
    assert ev.as_dict()['values']['colours.ink']['assumed'] is True
    assert 'black assumed' in ev.as_dict()['values']['colours.ink']['why']


def test_counts_a_colour_in_all_three_roles(filled_deck):
    deck = read_deck(filled_deck('#172033', '#56627A', '#F3F5F8', '#2350C8'))
    counts = census(deck)
    assert counts['#172033']['runs'] == 24     # two ink runs per slide, twelve slides
    assert counts['#F3F5F8']['fills'] == 12
    assert counts['#F3F5F8']['lines'] == 12


def test_names_ink_muted_panel_and_accent(filled_deck):
    ev = Evidence()
    out = colours(read_deck(filled_deck('#172033', '#56627A', '#F3F5F8', '#2350C8')), ev)
    assert out['ink'] == '#172033'
    assert out['muted'] == '#56627A'
    assert out['panel'] == '#F3F5F8'
    assert out['accent'] == '#2350C8'


def test_the_ink_is_the_most_used_text_colour_even_when_another_sorts_first(filled_deck):
    # The reference deck's ink (#172033) happens to sort first alphabetically, so the test above
    # passes even if role assignment falls back to sorting. Here #112233 sorts before #993333 and
    # is used half as often: USE has to decide.
    ev = Evidence()
    out = colours(read_deck(filled_deck('#993333', '#112233', '#F3F5F8', '#2350C8')), ev)
    assert out['ink'] == '#993333'
    assert out['muted'] == '#112233'


def test_a_colour_under_the_floor_is_incidental_not_a_token(filled_deck):
    # 12 slides paint the four tokens; a fifth colour appears on 3 shapes only.
    from io import BytesIO

    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    data = filled_deck('#172033', '#56627A', '#F3F5F8', '#2350C8')
    prs = Presentation(BytesIO(data))
    for slide in list(prs.slides)[:3]:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9), Inches(1),
                                     Inches(1), Inches(1))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string('BADA55')
    out = BytesIO()
    prs.save(out)

    ev = Evidence()
    tokens = colours(read_deck(out.getvalue()), ev)
    assert '#BADA55' not in tokens.values(), 'three uses is not a token'
    incidental = [i for i in ev.as_dict()['incidental'] if i['value'] == '#BADA55']
    assert incidental and 'under the 10-use floor' in incidental[0]['why']


def test_a_deck_whose_runs_state_no_colour_still_gets_an_ink(build_deck):
    # Review Focus 2: a deck whose runs carry no explicit colour. Without a fallback, colours.ink
    # is missing and the pack cannot validate.
    #
    # NOTE ON THE OLD NAME: this was called "resolves theme colours", which it never did — the
    # stock template's dk1 is a sysClr, so the assertion was satisfied by the black of last resort.
    # Resolving a scheme colour reference to RGB is NOT implemented (design §5 asks for it); the
    # two tests above cover the fallbacks that stand in for it, and it is an open item.
    from io import BytesIO

    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for _ in range(12):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(3), Inches(0.4))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = 'theme coloured'
        run.font.size = Pt(12)
    out = BytesIO()
    prs.save(out)

    ev = Evidence()
    tokens = colours(read_deck(out.getvalue()), ev)
    assert tokens.get('ink'), 'a deck with only theme colours still needs an ink'
    assert tokens.get('panel')
    assert ev.as_dict()['values']['colours.ink']['why'], 'the fallback has to say it is one'
