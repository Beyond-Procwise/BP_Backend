"""Findings from the whole-branch review of 2026-10-02, each with the failure it reproduces.

Kept in one file because they cut across modules; every one names the finding it closes.
"""
import json
import os
import shutil
import subprocess
import uuid

import pytest

from src.services.atb.pptx_import import store
from src.services.atb.pptx_import.contract import validate_layout, validate_pack
from src.services.atb.pptx_import.emit import layout_key
from src.services.atb.pptx_import.import_pack import import_pack

UI = os.environ.get('BEYOND_PROCWISE_UI', os.path.expanduser('~/PycharmProjects/beyond_procwise_ui'))
REFERENCE = os.environ.get(
    'ATB_REFERENCE_PACK',
    os.path.expanduser('~/Downloads/Infrastructure-Procurement-Strategy-Pack.pptx'))

def _pack_with_fonts(fonts: dict) -> dict:
    """A pack that is valid in every other respect, so only the fonts are under test."""
    return {
        'kind': 'atb_style', 'schema_version': 1, 'key': 'k', 'name': 'K',
        'format': {'kind': 'deck', 'width_in': 13.333, 'height_in': 7.5},
        'colours': {'ink': '#172033', 'muted': '#56627A', 'accent': '#12897F',
                    'panel': '#F4F7FA'},
        'type_scale_pt': {'title': 30, 'subtitle': 14, 'body': 11.5, 'table': 10, 'footer': 9},
        'fonts': fonts,
        'grid': {'cols': 12, 'margin_in': 0.5, 'gutter_in': 0.3, 'title_top_in': 0.35,
                 'body_top_in': 1.45},
        'series_palette': ['#12897F', '#0C6C9C'], 'rating_scales': {},
        'writing': {'locale': 'en-GB', 'locale_contested': False, 'locale_suggested': 'en-GB',
                    'title_max_words': 12, 'title_style': 'assertion',
                    'subtitle_style': 'basis'},
    }


SLIDE = [(0.5, 0.35, 12.33, 0.6, 'A reused title', 30, '#172033'),
         (0.5, 1.0, 12.33, 0.4, 'the basis', 14, '#56627A'),
         (0.5, 1.5, 6.0, 1.0, 'reused body text', 11.5, '#172033')]


# --------------------------------------------------------------------------- C2, in the browser
@pytest.mark.skipif(not os.path.exists(os.path.join(UI, 'src/modules/SpendIQ/atb/validateStyle.js'))
                    or not shutil.which('node') or not os.path.exists(REFERENCE),
                    reason='needs node, the UI checkout and ATB_REFERENCE_PACK')
def test_the_emitted_pack_and_layouts_pass_the_REAL_javascript_validators(tmp_path):
    """The drift test compares constant arrays. It stayed green while the pack was missing `kind`
    and `schema_version`, so every imported pack would have been rejected in the browser and
    silently replaced by the bundled ones. This runs the actual validators."""
    with open(REFERENCE, 'rb') as handle:
        result = import_pack(handle.read(), os.path.basename(REFERENCE), 'test')
    payload = tmp_path / 'emitted.json'
    payload.write_text(json.dumps({'pack': result.pack, 'layouts': result.layouts}))
    script = '''
import {validateStyle, validateLayoutAgainstStyle} from './src/modules/SpendIQ/atb/validateStyle.js';
import {validateLayout} from './src/modules/SpendIQ/atb/validateLayout.js';
import {readFileSync} from 'node:fs';
const {pack, layouts} = JSON.parse(readFileSync(process.env.ATB_PAYLOAD, 'utf8'));
const errors = validateStyle(pack).errors.slice();
for (const l of layouts) {
  errors.push(...validateLayout(l).errors, ...validateLayoutAgainstStyle(l, pack).errors);
}
console.log(JSON.stringify(errors));
'''
    done = subprocess.run(['node', '--input-type=module', '-e', script], cwd=UI,
                          capture_output=True, text=True, timeout=120,
                          env={**os.environ, 'ATB_PAYLOAD': str(payload)})
    assert done.returncode == 0, done.stderr[-400:]
    assert json.loads(done.stdout.strip().splitlines()[-1]) == []


# ------------------------------------------------- the box_in rule, on both sides of the fence
@pytest.mark.skipif(not os.path.exists(os.path.join(UI, 'src/modules/SpendIQ/atb/validateLayout.js'))
                    or not shutil.which('node'),
                    reason='needs node and the UI checkout')
def test_the_javascript_validator_enforces_the_box_in_rule_python_enforces(tmp_path):
    """contract.py's docstring said the JS validator "does not know about box_in yet". It does
    now, and this is what keeps the two from drifting: a region stating both is refused on both
    sides, and box_in alone is accepted on both. Without the JS half, every imported region
    positioned at the page origin in the browser and no validator said a word."""
    layout = {'id': 'x', 'version': 1, 'name': 'X', 'formats': ['deck'],
              'regions': [{'id': 'body', 'component': 'paragraph',
                           'box_in': {'x': 0.5, 'y': 1.4, 'w': 6, 'h': 2}}],
              'slots': {'body': {'type': 'text', 'fill': 'agent', 'max_chars': 100}}}
    both = json.loads(json.dumps(layout))
    both['regions'][0]['grid'] = {'col': 1, 'w': 6, 'row': 'body'}
    short = json.loads(json.dumps(layout))
    del short['regions'][0]['box_in']['h']

    payload = tmp_path / 'layouts.json'
    payload.write_text(json.dumps({'good': layout, 'both': both, 'short': short}))
    script = """
import {validateLayout} from './src/modules/SpendIQ/atb/validateLayout.js';
import {readFileSync} from 'node:fs';
const docs = JSON.parse(readFileSync(process.env.ATB_PAYLOAD, 'utf8'));
const out = {};
for (const [name, doc] of Object.entries(docs)) out[name] = validateLayout(doc).errors;
console.log(JSON.stringify(out));
"""
    done = subprocess.run(['node', '--input-type=module', '-e', script], cwd=UI,
                          capture_output=True, text=True, timeout=120,
                          env={**os.environ, 'ATB_PAYLOAD': str(payload)})
    assert done.returncode == 0, done.stderr[-400:]
    seen = json.loads(done.stdout.strip().splitlines()[-1])

    assert seen['good'] == []
    assert any('exactly one' in err for err in seen['both'])
    assert any('box_in needs numeric' in err for err in seen['short'])

    # and Python says the same three things about the same three documents
    assert validate_layout(layout) == []
    assert any('exactly one' in err for err in validate_layout(both))
    assert any('box_in needs numeric' in err for err in validate_layout(short))


# ------------------------------------------- a font family cannot carry a CSS declaration
@pytest.mark.skipif(not os.path.exists(os.path.join(UI, 'src/modules/SpendIQ/atb/validateStyle.js'))
                    or not shutil.which('node'),
                    reason='needs node and the UI checkout')
def test_both_validators_refuse_a_font_family_that_carries_css(tmp_path):
    """`fonts.*.family` is emitted into a CSS custom property inside a `style` attribute, and a
    pack is read out of a file somebody uploaded. Requiring only "a non-empty string" let a
    family close the value and append a declaration of its own — an outbound request to a host
    the uploader chose, fired by a reviewer merely looking at a candidate paper.

    Runs the REAL JavaScript validator and the vendored Python one over the same strings."""
    evil = 'A"; background-image:url(https://evil.example/x.png); --z:"'
    cases = {
        'evil_family': _pack_with_fonts({'heading': {'family': evil, 'fallback': 'sans-serif'},
                                         'body': {'family': 'Calibri', 'fallback': 'sans-serif'}}),
        'evil_fallback': _pack_with_fonts(
            {'heading': {'family': 'Cambria', 'fallback': 'serif; background:url(http://e/x)'},
             'body': {'family': 'Calibri', 'fallback': 'sans-serif'}}),
        'real': _pack_with_fonts(
            {'heading': {'family': 'Aptos Display', 'fallback': "Georgia, 'Times New Roman', serif"},
             'body': {'family': '微软雅黑',
                      'fallback': "'Segoe UI', system-ui, -apple-system, Arial, sans-serif"}}),
    }
    payload = tmp_path / 'packs.json'
    payload.write_text(json.dumps(cases))
    script = """
import {validateStyle} from './src/modules/SpendIQ/atb/validateStyle.js';
import {readFileSync} from 'node:fs';
const packs = JSON.parse(readFileSync(process.env.ATB_PAYLOAD, 'utf8'));
const out = {};
for (const [name, pack] of Object.entries(packs)) out[name] = validateStyle(pack).errors;
console.log(JSON.stringify(out));
"""
    done = subprocess.run(['node', '--input-type=module', '-e', script], cwd=UI,
                          capture_output=True, text=True, timeout=120,
                          env={**os.environ, 'ATB_PAYLOAD': str(payload)})
    assert done.returncode == 0, done.stderr[-400:]
    seen = json.loads(done.stdout.strip().splitlines()[-1])

    assert any('family' in err for err in seen['evil_family']), seen['evil_family']
    assert any('fallback' in err for err in seen['evil_fallback']), seen['evil_fallback']
    # A real deck's own families and stacks, including a non-ASCII family, still pass.
    assert seen['real'] == []

    # And Python says the same three things about the same three packs.
    assert any('family' in err for err in validate_pack(cases['evil_family']))
    assert any('fallback' in err for err in validate_pack(cases['evil_fallback']))
    assert validate_pack(cases['real']) == []


# ------------------------------------------------------------------------------------------- C5
@pytest.fixture
def conn():
    if os.environ.get('PROCWISE_TEST_LIVE_DB') != '1':
        pytest.skip('set PROCWISE_TEST_LIVE_DB=1')
    from dotenv import load_dotenv
    load_dotenv()
    from src.services.db import get_conn
    with get_conn() as connection:
        yield connection


@pytest.fixture
def name():
    return 'test-review-%s.pptx' % uuid.uuid4().hex[:10]


def test_a_crashed_import_does_not_erase_the_names_a_human_gave(conn, name, build_deck):
    """inherited() took the highest version that was not rejected — including one left at
    `importing` by a crash, which has no layouts and no names. Every name and every hand-defined
    scale from the last good version was then lost on the next import."""
    first = import_pack(build_deck([SLIDE] * 12), name, 'tester', conn=conn)
    layout_id = store.layouts(conn, pack_id=first.pack_id)[0]['layout_id']
    store.rename_layout(conn, layout_id, 'Eight headline recommendations', 'nick')
    store.define_rating_scale(conn, first.pack_id, 'hml',
                              {'High': {'bg': '#F9E1E1', 'ink': '#B42D2D'}}, 'nick')

    # a crash between insert_importing and mark_candidate
    store.insert_importing(conn, pack_key=store_key(name), version=2, source_file=name,
                           source_sha256='x', slide_count=12, pack=first.pack,
                           evidence={}, user='tester')

    carried = store.inherited(conn, store_key(name))
    assert carried['names'][first.layouts[0]['id']] == 'Eight headline recommendations'
    assert 'hml' in carried['rating_scales']


def store_key(filename):
    from src.services.atb.pptx_import.import_pack import key_for
    return key_for(filename)


# ------------------------------------------------------------------------------------------- I6
def test_an_importing_pack_cannot_be_approved_around_the_status_filter(conn, name, build_deck):
    """Every read path excludes `importing`. set_pack_status had no status guard, so one UPDATE
    walked an incomplete pack to `approved`, after which it and its missing layouts were served."""
    pack_id = store.insert_importing(conn, pack_key=store_key(name), version=1, source_file=name,
                                     source_sha256='x', slide_count=1,
                                     pack={'key': 'k'}, evidence={}, user='tester')
    changed = store.set_pack_status(conn, pack_id, 'approved', 'nick')
    assert changed == 0, 'an importing pack is not approvable'
    assert all(p['pack_id'] != pack_id for p in store.packs(conn))


def test_approving_something_that_does_not_exist_changes_nothing(conn):
    assert store.set_pack_status(conn, '00000000-0000-0000-0000-000000000000', 'approved', 'n') == 0
    assert store.set_layout_status(conn, '00000000-0000-0000-0000-000000000000', 'approved', 'n') == 0
    assert store.rename_layout(conn, '00000000-0000-0000-0000-000000000000', 'x', 'n') == 0


# ------------------------------------------------------------------------------------------ I12
def test_two_packs_with_the_same_structure_get_different_layout_ids(build_deck):
    """layout_key hashed only the signature, so every pack with a full-width table emitted the
    same id. bridge.js keys its registry by id: two approved packs would offer one picker entry
    and render the other pack's geometry."""
    first = import_pack(build_deck([SLIDE] * 12), 'alpha.pptx', 'tester')
    second = import_pack(build_deck([SLIDE] * 12), 'beta.pptx', 'tester')
    assert first.layouts[0]['id'] != second.layouts[0]['id']
    # ...and still stable for the same pack
    again = import_pack(build_deck([SLIDE] * 12), 'alpha.pptx', 'tester')
    assert again.layouts[0]['id'] == first.layouts[0]['id']


def test_the_layout_key_is_still_lower_snake_case(build_deck):
    import re
    first = import_pack(build_deck([SLIDE] * 12), 'Alpha Pack 2026.pptx', 'tester')
    assert re.match(r'^[a-z][a-z0-9_]*$', first.layouts[0]['id'])


# ------------------------------------------------------------------------------------------ I11
def test_a_smaller_16_9_deck_keeps_its_body_row(build_deck):
    """PowerPoint's other 16:9 size is 10 x 5.625in. With an absolute 1.25in title band the body's
    lead row at 1.125in fell INSIDE the band and vanished from the layout, with no problem."""
    from src.services.atb.pptx_import.cluster import group
    from src.services.atb.pptx_import.read import read_deck

    slide = [(0.375, 0.26, 9.25, 0.45, 'A title', 22, '#172033'),
             (0.375, 1.125, 4.5, 0.75, 'the lead row', 9, '#172033'),
             (0.375, 2.25, 2.1, 1.5, 'a', 9, '#172033'),
             (2.7, 2.25, 2.1, 1.5, 'b', 9, '#172033'),
             (5.0, 2.25, 2.1, 1.5, 'c', 9, '#172033')]
    deck = read_deck(build_deck([slide] * 8, width_in=10.0, height_in=5.625))
    signature = group(deck)[0].signature
    assert len(signature) == 2, 'the lead row and the three cards are two rows'
    assert signature[0][0] == 1 and signature[1][0] == 3


def test_a_full_bleed_background_does_not_become_the_title_line(build_deck):
    """A background rectangle is in every real deck. Counted, 12 shapes at y=0 beat 12 at y=0.35
    and the title band came back at 0.0in."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.grid import grid
    from src.services.atb.pptx_import.read import read_deck

    slide = [(0.0, 0.0, 13.333, 7.5, '', 12, '#F3F5F8'),
             (0.5, 0.35, 12.33, 0.6, 'A title', 30, '#172033'),
             (0.5, 1.5, 6.0, 1.0, 'body', 11.5, '#172033')]
    g = grid(read_deck(build_deck([slide] * 12)), Evidence())
    assert g['title_top_in'] == 0.35
    assert g['margin_in'] == 0.5


def test_the_body_top_is_the_first_row_not_the_most_crowded_one(build_deck):
    """The modal mid-band edge lands on the biggest card row. Here: one lead row at 1.5in and four
    cards at 3.0in — the body starts at 1.5."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.grid import grid
    from src.services.atb.pptx_import.read import read_deck

    slide = [(0.5, 0.35, 12.33, 0.6, 'A title', 30, '#172033'),
             (0.5, 1.5, 12.33, 0.5, 'the lead row', 11.5, '#172033')] + \
            [(0.5 + i * 3.2, 3.0, 2.9, 2.0, 'card %d' % i, 11.5, '#172033') for i in range(4)]
    assert grid(read_deck(build_deck([slide] * 10)), Evidence())['body_top_in'] == 1.5


# ------------------------------------------------------------------- C3: an assumed pack SAYS so
def test_a_deck_that_states_no_colour_reports_every_assumption_as_a_problem(build_deck):
    """A theme-driven deck produced ink/muted/panel/accent entirely from fallbacks, with
    `problems: []` and evidence that read like a measurement. A pack of assumptions must say so
    where a human looks."""
    from io import BytesIO

    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for _ in range(12):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for y, text in ((0.35, 'A title'), (1.5, 'body text')):
            box = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(6), Inches(0.5))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = text
            run.font.size = Pt(30 if y < 1 else 11)
    out = BytesIO()
    presentation.save(out)

    result = import_pack(out.getvalue(), 'themed.pptx', 'tester')
    assumed = [p for p in result.problems if p['kind'] == 'assumed']
    paths = {p['region'] for p in assumed}
    assert 'colours.ink' in paths
    assert 'fonts.heading' in paths, 'inherit is not a measured family'
    assert all(p['why'] for p in assumed)


# ------------------------------------------------------------------------- C4, I8, I9 and I13
PACK = {'grid': {'cols': 12, 'margin_in': 0.5, 'gutter_in': 0.3, 'title_top_in': 0.35,
                 'body_top_in': 1.5, 'footer_top_in': 7.02},
        'type_scale_pt': {'title': 30, 'subtitle': 14, 'card_title': 13, 'body': 11.5,
                          'table': 10, 'footer': 9},
        'format': {'kind': 'deck', 'width_in': 13.333, 'height_in': 7.5}}


def _shape(x, y, w, h, kind='text', slide=1, words='', table=None):
    from src.services.atb.pptx_import.read import Box, Run, Shape
    runs = (Run(text=words, size_pt=12, font='Calibri', colour='#172033', bold=False,
                lang='en-GB'),) if words else ()
    return Shape(kind=kind, box=Box(x, y, w, h), runs=runs, fill=None, line=None,
                 table=table, slide=slide, name='s')


def _deck(slides):
    from src.services.atb.pptx_import.read import Deck
    return Deck(width_in=13.333, height_in=7.5, slides=tuple(slides),
                theme={'colours': {}, 'fonts': {}}, chart_series_colours=(), run_langs={})


def _cluster(signature, rows, slides=(1, 2), members=None):
    from src.services.atb.pptx_import.cluster import Cluster
    return Cluster(signature=signature, slides=slides, rows=rows,
                   rows_by_slide=tuple(members or (rows, rows)))


def test_a_tall_row_is_clipped_to_the_next_row_instead_of_swallowing_it():
    """A row's span is the union of its shapes, so one tall panel beside small cards produced a
    region that covered every row below it — 7 of the reference deck's 8 layouts overlapped."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.slots import regions_and_slots

    tall = (_shape(0.5, 1.5, 3.0, 4.0, words='a side panel'),
            _shape(4.0, 1.5, 8.0, 0.8, words='first card'))
    below = (_shape(4.0, 3.0, 8.0, 0.8, words='second card'),)
    cluster = _cluster(((2, 'SHAPE'), (1, 'SHAPE')), (tall, below))
    regions, _, problems = regions_and_slots(cluster, _deck([tall + below]), PACK, Evidence())
    body = [r for r in regions if 'box_in' in r and r['id'] != 'title']
    first, second = body[0]['box_in'], body[1]['box_in']
    assert first['y'] + first['h'] <= second['y'], 'a band ends where the next begins'
    assert not [p for p in problems if p['kind'] == 'overlap']


def test_a_region_that_runs_into_the_footer_is_reported():
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.slots import regions_and_slots

    deep = (_shape(0.5, 5.0, 12.33, 2.5, words='a very deep row'),)
    cluster = _cluster(((1, 'SHAPE'),), (deep,))
    _, _, problems = regions_and_slots(cluster, _deck([deep]), PACK, Evidence())
    footer = [p for p in problems if p['kind'] == 'into_footer']
    assert footer and '7.02' in footer[0]['why']


def test_a_region_past_the_right_margin_is_reported():
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.slots import regions_and_slots

    wide = (_shape(0.5, 1.5, 13.0, 1.0, words='too wide'),)
    cluster = _cluster(((1, 'SHAPE'),), (wide,))
    _, _, problems = regions_and_slots(cluster, _deck([wide]), PACK, Evidence())
    assert [p['kind'] for p in problems if p['kind'] == 'off_page'] == ['off_page']


def test_a_repeating_row_is_one_region_that_says_it_repeats():
    """The signature collapses a trailing repeat; iterating the raw rows emitted six card regions
    copied from whichever slide came first, which no renderer can know to repeat."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.slots import regions_and_slots

    band = lambda y: tuple(_shape(0.5 + i * 3.2, y, 2.9, 1.0, words='card') for i in range(4))
    rows = (band(1.5), band(2.8), band(4.1))
    cluster = _cluster(((4, 'SHAPE', 'repeat'),), rows)
    regions, slots, _ = regions_and_slots(cluster, _deck([sum(rows, ())]), PACK, Evidence())
    cards = [r for r in regions if r['component'] == 'text_card']
    assert len(cards) == 1, 'three identical bands are one repeating region'
    assert cards[0]['repeat_over'] == cards[0]['id']
    assert slots[cards[0]['id']]['max'] == 12


def test_a_table_whose_height_varies_is_not_a_loose_fit():
    """A table's height is its row count. Comparing it reported `slide 4 is 3.8in off` and made 41
    of the reference deck's 45 problems artefacts."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.slots import regions_and_slots

    rows = (('A', 'B'), ('c', 'd'))
    short = ((_shape(0.5, 1.5, 12.33, 1.0, kind='table', slide=1, table=rows),),)
    deep = ((_shape(0.5, 1.5, 12.33, 4.5, kind='table', slide=2, table=rows),),)
    cluster = _cluster(((1, 'TABLE'),), short[0] and (short[0][0],) and short,
                       members=(short, deep))
    regions, _, problems = regions_and_slots(cluster, _deck([short[0], deep[0]]), PACK, Evidence())
    assert not [p for p in problems if p['kind'] == 'loose_fit']
    table = [r for r in regions if r['component'] == 'table'][0]
    assert table['height_range_in'] == [1.0, 4.5], 'the range is recorded, not flagged'


def test_the_loose_fit_tolerance_is_the_thing_being_tested():
    """`LOOSE_FIT_IN = 0.0` left the whole suite green: every fixture had one member, or two
    identical ones, so `off` was exactly zero whatever the tolerance."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.slots import LOOSE_FIT_IN, regions_and_slots

    def problems_for(second_width):
        first = ((_shape(0.5, 1.5, 6.0, 1.0, slide=1, words='x'),),)
        second = ((_shape(0.5, 1.5, second_width, 1.0, slide=2, words='x'),),)
        cluster = _cluster(((1, 'SHAPE'),), first, members=(first, first, second))
        _, _, found = regions_and_slots(cluster, _deck([first[0], second[0]]), PACK, Evidence())
        return [p for p in found if p['kind'] == 'loose_fit']

    assert LOOSE_FIT_IN == 0.15
    assert problems_for(6.0 + 0.14) == [], 'inside the tolerance'
    assert problems_for(6.0 + 0.16), 'outside it'


def test_each_region_is_filled_from_its_own_row():
    """Every text slot got the same sentence and every card grid the same leading items, so the
    thumbnail a human approves already differed from the page the report prints."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.example import example_fill
    from src.services.atb.pptx_import.slots import regions_and_slots

    first = (_shape(0.5, 1.5, 12.33, 0.6, words='the first paragraph'),)
    second = (_shape(0.5, 3.0, 12.33, 0.6, words='a different second paragraph'),)
    cluster = _cluster(((1, 'SHAPE'), (1, 'SHAPE')), (first, second))
    _, slots, _ = regions_and_slots(cluster, _deck([first + second]), PACK, Evidence())
    fill, _ = example_fill(cluster, _deck([first + second]), slots)
    texts = [slot['text'] for name, slot in fill['slots'].items() if name.startswith('prose')]
    assert texts == ['the first paragraph', 'a different second paragraph']


def test_a_mid_tone_fill_is_not_a_panel():
    """`_PANEL_LUMINANCE 0.85 -> 0.5` left the suite green: no fixture had a mid-tone fill."""
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.palette import colours

    mid = tuple((_shape(0.5, 2.0, 3.0, 0.4, slide=i),) for i in range(1, 13))
    deck = _deck(mid)
    shapes = []
    for i in range(1, 13):
        from src.services.atb.pptx_import.read import Box, Run, Shape
        shapes.append((Shape(kind='auto', box=Box(0.5, 2.0, 3, 0.4), runs=(
            Run(text='x', size_pt=12, font='Calibri', colour='#172033', bold=False, lang='en'),),
            fill='#7A8A99', line=None, table=None, slide=i, name='s'),))
    out = colours(_deck(tuple(shapes)), Evidence())
    assert out['panel'] != '#7A8A99', 'a mid-tone fill is not the pale panel'


def test_an_american_deck_with_british_words_is_decided_by_the_comparison():
    """`AMERICAN_RE` could be neutered with the suite green: the American deck had zero British
    words, so the count comparison never decided anything."""
    from src.services.atb.pptx_import.derived import writing
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.read import Box, Deck, Run, Shape

    def deck_of(words):
        runs = tuple(Run(text=w, size_pt=12, font='Calibri', colour='#172033', bold=False,
                         lang='en-US') for w in words)
        shape = Shape(kind='text', box=Box(0.5, 1.5, 6, 1), runs=runs, fill=None, line=None,
                      table=None, slide=1, name='t')
        return Deck(width_in=13.333, height_in=7.5, slides=((shape,),),
                    theme={'colours': {}, 'fonts': {}}, chart_series_colours=(),
                    run_langs={'en-US': len(words)})

    mostly_american = ['organize', 'organized', 'organizing', 'optimization', 'realized',
                       'colorize', 'itemized', 'Virtualisation', 'Mobilise', 'Optimise']
    assert writing(deck_of(mostly_american), Evidence())['locale_contested'] is False
    mostly_british = ['Virtualisation', 'Mobilise', 'Optimise', 'utilisation', 'organize']
    assert writing(deck_of(mostly_british), Evidence())['locale_contested'] is True


def test_a_genuine_british_spelling_counts_as_one():
    """labour, colour, behaviour and the rest were in the false-friends list — the British forms
    themselves — so a deck writing them fifty times scored zero and was never contested."""
    from src.services.atb.pptx_import.derived import writing
    from src.services.atb.pptx_import.evidence import Evidence
    from src.services.atb.pptx_import.read import Box, Deck, Run, Shape

    words = ['colour', 'behaviour', 'labour', 'favour', 'honour', 'rigour']
    runs = tuple(Run(text=w, size_pt=12, font='Calibri', colour='#172033', bold=False,
                     lang='en-US') for w in words)
    shape = Shape(kind='text', box=Box(0.5, 1.5, 6, 1), runs=runs, fill=None, line=None,
                  table=None, slide=1, name='t')
    deck = Deck(width_in=13.333, height_in=7.5, slides=((shape,),),
                theme={'colours': {}, 'fonts': {}}, chart_series_colours=(),
                run_langs={'en-US': len(words)})
    result = writing(deck, Evidence())
    assert result['locale_contested'] is True
    assert result['locale_suggested'] == 'en-GB'
