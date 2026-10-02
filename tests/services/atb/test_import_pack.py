import pytest

from src.services.atb.pptx_import.import_pack import ImportRefused, import_pack

REUSED = [(0.5, 0.35, 12.33, 0.6, 'A reused title', 30, '#172033'),
          (0.5, 1.0, 12.33, 0.4, 'the basis for it', 14, '#56627A'),
          (0.5, 1.5, 6.0, 1.0, 'reused body text', 11.5, '#172033')]
ONE_OFF = [(0.5, 0.35, 12.33, 0.6, 'A one-off page', 30, '#172033'),
           (0.5, 1.5, 2.0, 1.0, 'a', 11.5, '#172033'),
           (3.0, 1.5, 2.0, 1.0, 'b', 11.5, '#172033'),
           (5.5, 1.5, 2.0, 1.0, 'c', 11.5, '#172033')]


def test_emits_the_reused_layouts_and_lists_the_single_use_ones(build_deck):
    result = import_pack(build_deck([REUSED] * 11 + [ONE_OFF]), 'Sample.pptx', 'tester')
    assert result.layouts, 'the reused structure earns a layout'
    assert all(len(l['slide_refs']) > 1 for l in result.layouts)
    assert result.single_use, 'the one-off is listed, not dropped (design §6a)'
    assert result.single_use[0]['slides'] == [12]
    assert '3-up' in result.single_use[0]['structure']


def test_the_pack_key_comes_from_the_filename(build_deck):
    result = import_pack(build_deck([REUSED] * 12),
                         'Infrastructure-Procurement-Strategy-Pack.pptx', 'tester')
    assert result.pack_key == 'infrastructure-procurement-strategy-pack'
    assert result.pack['key'] == result.pack_key


def test_refuses_an_unreadable_file():
    with pytest.raises(ImportRefused) as exc:
        import_pack(b'nope', 'x.pptx', 'tester')
    assert 'could not be opened' in str(exc.value)


def test_refuses_a_deck_with_no_slides():
    from io import BytesIO

    from pptx import Presentation
    out = BytesIO()
    Presentation().save(out)
    with pytest.raises(ImportRefused) as exc:
        import_pack(out.getvalue(), 'x.pptx', 'tester')
    assert 'no slides' in str(exc.value)


def test_the_same_file_twice_gives_identical_output(build_deck):
    data = build_deck([REUSED] * 12)
    first = import_pack(data, 'Sample.pptx', 'tester')
    second = import_pack(data, 'Sample.pptx', 'tester')
    assert first.pack == second.pack
    assert first.layouts == second.layouts
    assert first.single_use == second.single_use


def test_the_evidence_comes_back_with_the_result(build_deck):
    result = import_pack(build_deck([REUSED] * 12), 'Sample.pptx', 'tester')
    assert result.evidence['values']['colours.ink']['value'] == '#172033'
    assert 'theme.colours' in result.evidence['ignored']


def test_problems_from_every_layout_are_collected(build_deck):
    # A chart row declares that it cannot draw yet; the result carries that up.
    from io import BytesIO

    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for _ in range(12):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.33), Inches(0.6))
        run = box.text_frame.paragraphs[0].add_run()
        run.text, run.font.size = 'A chart page', Pt(30)
        data = CategoryChartData()
        data.categories = ['a', 'b']
        data.add_series('s', (1.0, 2.0))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5),
                               Inches(6), Inches(3), data)
    out = BytesIO()
    prs.save(out)

    result = import_pack(out.getvalue(), 'Charts.pptx', 'tester')
    assert any(p['kind'] == 'unresolved' for p in result.problems)
    assert any('no chart renderer' in p['why'] for p in result.problems)


def test_nothing_is_stored_when_no_connection_is_given(build_deck):
    result = import_pack(build_deck([REUSED] * 12), 'Sample.pptx', 'tester')
    assert result.pack_id is None
    assert result.version == 1
