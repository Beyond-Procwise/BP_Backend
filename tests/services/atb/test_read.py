import pytest

from src.services.atb.pptx_import.read import DeckUnreadable, read_deck


def test_reads_page_size_and_one_shape(build_deck):
    deck = read_deck(build_deck([[(0.5, 0.35, 12.33, 0.6, 'A title', 30, '#172033')]]))
    assert (round(deck.width_in, 3), round(deck.height_in, 3)) == (13.333, 7.5)
    assert len(deck.slides) == 1
    shape = deck.slides[0][0]
    assert shape.kind == 'text'
    assert (shape.box.x, shape.box.y, shape.box.w) == (0.5, 0.35, 12.33)
    assert shape.runs[0].text == 'A title'
    assert shape.runs[0].size_pt == 30.0
    assert shape.runs[0].colour == '#172033'
    assert shape.slide == 1


def test_flattens_groups_or_the_census_sees_nothing(grouped_deck):
    # Review Focus 1: PowerPoint groups freely. Unflattened, these two shapes are invisible.
    deck = read_deck(grouped_deck)
    texts = [r.text for s in deck.slides for sh in s for r in sh.runs]
    assert 'inside a group' in texts
    assert 'also inside' in texts
    assert all(sh.kind != 'group' for s in deck.slides for sh in s)


def test_refuses_a_file_it_cannot_open():
    with pytest.raises(DeckUnreadable) as e:
        read_deck(b'this is not a pptx')
    assert 'could not be opened' in str(e.value)


def test_refuses_a_deck_with_no_slides():
    from io import BytesIO

    from pptx import Presentation
    out = BytesIO()
    Presentation().save(out)
    with pytest.raises(DeckUnreadable) as e:
        read_deck(out.getvalue())
    assert 'no slides' in str(e.value)
