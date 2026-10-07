"""The doors: who may use presentation data, what an export carries, who may grant scope."""
import datetime as dt
import io
import time
from types import SimpleNamespace as NS

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from api.routers import reports as rr
from src.services import rbac, report_export
from src.services.report_data import live, mode as rd_mode, service
from src.services.report_data.live import Row
from src.services.report_data.scope import Scope

NOW = int(time.time())


def principal(role="Admin", sub="u-admin", jti="j1", exp=None):
    return NS(subject=sub, email=f"{sub}@x", claims={"jti": jti, "auth_time": NOW, "exp": exp or NOW + 3600, "_role": role})


class Log:
    """Stands in for proc.bp_presentation_log and the scope table."""
    def __init__(self):
        self.rows, self.scope = [], []


@pytest.fixture
def env(monkeypatch):
    st = NS(role="Admin", who=principal(), log=Log(), denied=set(), rendered={}, stored=[], exports=[])
    monkeypatch.setattr(rbac, "effective_role", lambda p, *a, **k: p.claims["_role"])

    def _log(subject, sid, event, detail=None):
        st.log.rows.append((subject, sid, event, detail or {}))
    monkeypatch.setattr(rd_mode, "_log", _log)

    class Cur:
        def __init__(self): self.out = None
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def execute(self, sql, params=None):
            if "FROM proc.bp_presentation_log" in sql:
                subj, sid = params
                ev = [r[2] for r in st.log.rows if r[0] == subj and r[1] == sid and r[2] in ("activate", "deactivate")]
                self.out = [(ev[-1],)] if ev else []
        def fetchone(self): return self.out[0] if self.out else None
    class Conn:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def cursor(self): return Cur()
    import contextlib
    import src.services.db as dbm
    monkeypatch.setattr(dbm, "get_conn", lambda: Conn())

    def gate(action, principal_, **kw):
        if action in st.denied or (action == "report.scope.write" and principal_.claims["_role"] != "Admin"):
            raise HTTPException(status_code=403, detail="refused")
    monkeypatch.setattr(rr, "gate", gate)
    monkeypatch.setattr(live, "series", lambda m, g, *a: ([Row((dt.date(2026, 1, 1),), ("Jan 2026",), 5.0)] if g else [Row((), (), 5.0)]))
    monkeypatch.setattr(live, "unattributed_findings", lambda *a, **k: 0)
    monkeypatch.setattr(rd_service_scope(), "resolve", lambda p: Scope("Admin", True, ()) if p.claims["_role"] == "Admin" else Scope("Buyer", False, ("CC1",)))

    def render_pdf(html, css=""):
        st.rendered.update(html=html, css=css)
        return b"%PDF-fake"
    monkeypatch.setattr(report_export, "render_pdf", render_pdf)
    monkeypatch.setattr(report_export, "store", lambda c, k, t: st.stored.append(k) or True)
    monkeypatch.setattr(report_export, "record_url", lambda *a: True)
    app = FastAPI(); app.include_router(rr.router)
    app.dependency_overrides[rr.require_user] = lambda: st.who
    st.client = TestClient(app)
    return st


def rd_service_scope():
    import src.services.report_data.scope as s
    return s


REQ = {"from": "2026-01-01", "to": "2026-03-31", "tiles": [{"id": "a", "metric": "committed_spend", "viz": "kpi"}]}
PRES = {**REQ, "data_mode": "presentation"}


def test_the_default_is_live_and_a_tile_says_so(env):
    r = env.client.post("/reports/data", json=REQ)
    assert r.status_code == 200 and r.json()["data_mode"] == "live" and r.json()["tiles"][0]["data_mode"] == "live"
    assert r.json()["tiles"][0]["marker"] is None


@pytest.mark.parametrize("role", ["Buyer", "Viewer", "Approver"])
def test_a_non_admin_is_refused_presentation_data_on_data_export_and_even_a_handedited_request(env, role):
    env.who = principal(role, "u-b")
    env.client.app.dependency_overrides[rr.require_user] = lambda: env.who
    # a non-admin cannot switch it on ...
    assert env.client.post("/reports/presentation-mode", json={"active": True}).status_code == 403
    # ... and a hand-edited request cannot reach it, even if an activation row existed for the session
    env.log.rows.append((env.who.subject, rd_mode.session_id(env.who), "activate", {}))
    assert env.client.post("/reports/data", json=PRES).status_code == 403
    assert env.client.post("/reports/export", json={"format": "xlsx", "request": PRES}).status_code == 403
    assert env.client.post("/reports/export", json={"format": "pdf", "html": "<p>x</p>", "request": PRES}).status_code == 403
    assert env.stored == []


def test_an_admin_needs_an_active_activation_on_this_session(env):
    assert env.client.post("/reports/data", json=PRES).status_code == 403          # admin, but not switched on
    assert env.client.get("/reports/presentation-mode").json() == {"eligible": True, "active": False}
    assert env.client.post("/reports/presentation-mode", json={"active": True}).status_code == 200
    r = env.client.post("/reports/data", json=PRES)
    assert r.status_code == 200 and r.json()["marker"] == "PRESENTATION DATA - NOT REAL"
    assert r.json()["tiles"][0]["marker"] == "PRESENTATION DATA - NOT REAL"
    env.client.post("/reports/presentation-mode", json={"active": False})
    assert env.client.post("/reports/data", json=PRES).status_code == 403          # off again, nothing reachable


def test_presentation_ends_with_the_session(env):
    env.client.post("/reports/presentation-mode", json={"active": True})
    assert rd_mode.is_active(env.who)
    # a new sign-in is a new session: the old activation does not follow the user
    env.who = principal(jti="j2")
    env.client.app.dependency_overrides[rr.require_user] = lambda: env.who
    assert env.client.post("/reports/data", json=PRES).status_code == 403
    # an expired token is never active, activation row or not
    old = principal(exp=NOW - 10)
    env.log.rows.append((old.subject, rd_mode.session_id(old), "activate", {}))
    assert not rd_mode.is_active(old)
    with pytest.raises(PermissionError):
        rd_mode.activate(old)


def test_every_activation_and_presentation_export_is_logged(env):
    env.client.post("/reports/presentation-mode", json={"active": True})
    env.client.post("/reports/export", json={"format": "xlsx", "name": "Board", "request": PRES})
    events = [r[2] for r in env.log.rows]
    assert events == ["activate", "export"] and env.log.rows[0][0] == "u-admin"


def test_a_presentation_pdf_is_watermarked_on_every_page_and_never_stored(env):
    env.client.post("/reports/presentation-mode", json={"active": True})
    r = env.client.post("/reports/export", json={"format": "pdf", "name": "Board", "report_id": "7",
                                                 "html": "<html><body><p>layout</p></body></html>", "request": PRES})
    assert r.status_code == 200 and r.headers["x-report-data-mode"] == "presentation"
    assert "rb-watermark" in env.rendered["html"] and "PRESENTATION DATA - NOT REAL" in env.rendered["html"]
    assert "@page" in env.rendered["css"] and "PRESENTATION DATA - NOT REAL" in env.rendered["css"]   # margin boxes repeat per page
    assert env.stored == [] and r.headers["x-report-stored"] == "false"                              # contained


def test_a_presentation_workbook_is_flagged_on_the_first_sheet_and_in_the_filename(env):
    from openpyxl import load_workbook
    env.client.post("/reports/presentation-mode", json={"active": True})
    r = env.client.post("/reports/export", json={"format": "xlsx", "name": "Board", "request": PRES})
    assert 'filename="Board-DEMO.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    first = wb[wb.sheetnames[0]]
    assert first["A1"].value == "NOTE" and first["B1"].value == "PRESENTATION DATA - NOT REAL"
    assert wb[wb.sheetnames[1]]["A1"].value == "PRESENTATION DATA - NOT REAL"                         # and on the tile's own sheet


def test_a_live_export_carries_every_tile_in_its_true_state_with_its_parameters(env):
    from openpyxl import load_workbook
    env.denied = {"finding.read"}
    req = {**REQ, "tiles": [
        {"id": "spend", "metric": "committed_spend", "groupBy": ["month"], "viz": "line", "filters": {"region": ["North"]}},
        {"id": "dups", "metric": "duplicate_risk", "viz": "kpi"},                    # forbidden for this caller
        {"id": "tail", "metric": "compliance_rate", "viz": "kpi"},                   # no live source
        {"id": "bad", "metric": "committed_spend", "groupBy": ["nope"], "viz": "table"}]}
    r = env.client.post("/reports/export", json={"format": "pdf", "name": "Q1", "html": "<html><body>layout</body></html>", "request": req})
    html = env.rendered["html"]
    for tid in ("spend", "dups", "tail", "bad"):
        assert f'data-tile="{tid}"' in html                                          # nothing omitted
    assert "forbidden" in html and "no data source" in html and "rejected" in html
    assert "Grouped by: month" in html and "North" in html and "Period: 2026-01-01 to 2026-03-31" in html
    assert "as of" in html and "CURRENT_DATE" in html and "rb-watermark" not in html   # live: no mark
    x = env.client.post("/reports/export", json={"format": "xlsx", "name": "Q1", "request": req})
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Parameters", "spend", "dups", "tail", "bad"]
    assert 'Q1.xlsx' in x.headers["content-disposition"] and "-DEMO" not in x.headers["content-disposition"]


def test_an_export_is_computed_under_the_exporters_rights_not_the_authors(env):
    env.who = principal("Buyer", "u-b")
    env.client.app.dependency_overrides[rr.require_user] = lambda: env.who
    env.denied = {"invoice.read"}
    env.client.post("/reports/export", json={"format": "pdf", "name": "Q1", "html": "<html><body>x</body></html>", "request": REQ})
    assert "forbidden" in env.rendered["html"]


def test_scope_is_granted_and_revoked_only_by_an_admin_and_never_to_oneself(env):
    env.who = principal("Buyer", "u-b")
    env.client.app.dependency_overrides[rr.require_user] = lambda: env.who
    assert env.client.post("/reports/scope", json={"subject": "u-b", "buyer_id": "CC1"}).status_code == 403
    assert env.client.get("/reports/scope").status_code == 403
    env.who = principal("Admin", "u-admin")
    env.client.app.dependency_overrides[rr.require_user] = lambda: env.who
    assert env.client.post("/reports/scope", json={"subject": "u-admin", "buyer_id": "CC1"}).status_code == 403


def test_the_picker_lists_only_what_the_caller_may_read_and_hides_presentation_only_tiles_in_live_mode(env):
    keys = lambda r: {m["key"] for m in r.json()["metrics"]}
    live_keys = keys(env.client.get("/reports/registry"))
    assert "committed_spend" in live_keys and "compliance_rate" not in live_keys and "spend_by_category" not in live_keys
    env.denied = {"invoice.read"}
    assert "committed_spend" not in keys(env.client.get("/reports/registry"))      # no access, not listed
    assert "quote_volume" in keys(env.client.get("/reports/registry"))
    env.denied = set()
    env.client.post("/reports/presentation-mode", json={"active": True})
    r = env.client.get("/reports/registry")
    assert "compliance_rate" in keys(r) and {m["mode"] for m in r.json()["metrics"] if m["key"] == "compliance_rate"} == {"presentation"}
    assert "category" in {d["key"] for d in r.json()["dimensions"]}
    env.client.post("/reports/presentation-mode", json={"active": False})
    assert "category" not in {d["key"] for d in env.client.get("/reports/registry").json()["dimensions"]}


def test_a_new_metric_is_available_by_registry_entry_alone(env, monkeypatch):
    from src.services.report_data import registry as R
    extra = R.Metric("po_value_test", "PO value (test)", "currency", R.LIVE, "invoice", "SUM(i.invoice_amount)",
                     additive=True, dimensions=("month",))
    monkeypatch.setitem(R.METRICS, "po_value_test", extra)
    assert "po_value_test" in {m["key"] for m in env.client.get("/reports/registry").json()["metrics"]}
    out = env.client.post("/reports/data", json={**REQ, "tiles": [{"id": "n", "metric": "po_value_test", "viz": "kpi"}]})
    assert out.json()["tiles"][0]["status"] == "ok"                                 # no code outside the registry changed
    bad = env.client.post("/reports/data", json={**REQ, "tiles": [{"id": "n", "metric": "po_value_test", "groupBy": ["supplier"]}]})
    assert bad.json()["tiles"][0]["status"] == "rejected"                           # and it only allows what it lists


def test_filter_values_are_scoped_gated_and_never_leak_through_a_refusal(env, monkeypatch):
    seen = {}
    monkeypatch.setattr(live, "values", lambda src, dim, q, scope, limit=50: seen.update(src=src, dim=dim, q=q, scope=scope) or [("S1", "Supplier One")])
    body = {"metric": "committed_spend", "dimension": "supplier", "q": " sup "}
    r = env.client.post("/reports/values", json=body)
    assert r.status_code == 200 and r.json()["values"] == [{"key": "S1", "label": "Supplier One"}]
    assert seen["q"] == "sup" and seen["src"] == "invoice" and seen["scope"].all_rows
    env.who = principal("Buyer", "u-b")
    env.client.app.dependency_overrides[rr.require_user] = lambda: env.who
    env.client.post("/reports/values", json=body)
    assert not seen["scope"].all_rows and seen["scope"].buyers == ("CC1",)          # a Buyer is offered only their own
    env.denied = {"invoice.read"}
    r = env.client.post("/reports/values", json=body)
    assert r.json() == {"data_mode": "live", "status": "forbidden", "values": []}      # a refusal, not an empty list
    env.denied = set()
    assert env.client.post("/reports/values", json={**body, "dimension": "month"}).status_code == 422
    assert env.client.post("/reports/values", json={**body, "dimension": "detector_type"}).status_code == 422
    assert env.client.post("/reports/values", json={**body, "metric": "nope"}).status_code == 422


def test_presentation_values_are_synthetic_and_admin_only(env):
    body = {"metric": "committed_spend", "dimension": "supplier", "data_mode": "presentation"}
    assert env.client.post("/reports/values", json=body).status_code == 403
    env.client.post("/reports/presentation-mode", json={"active": True})
    r = env.client.post("/reports/values", json=body)
    labels = [v["label"] for v in r.json()["values"]]
    assert r.status_code == 200 and labels and all(l.startswith("Synthetic Supplier") for l in labels)
    assert r.json()["marker"] == "PRESENTATION DATA - NOT REAL"
    assert env.client.post("/reports/values", json={**body, "dimension": "finding_type"}).status_code == 422
