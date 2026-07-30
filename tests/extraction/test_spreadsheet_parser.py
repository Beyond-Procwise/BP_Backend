"""Spreadsheet (.xlsx/.csv) parsing → ParsedDocument, and end-to-end extraction."""
import os

import pytest

from src.services.extraction_v3.parsers.router import parse as route_parse
from src.services.extraction_v3.parsers.spreadsheet_backend import parse_csv, parse_xlsx

XLSX = "/home/muthu/Downloads/new/HR Quote v2.xlsx"


def test_csv_parses_to_markdown(tmp_path):
    f = tmp_path / "q.csv"
    f.write_text("Quote ID,Amount\nQTE-2026-01521,12000\nItem,Qty\nHR Programme,1\n")
    doc = parse_csv(f)
    assert doc.file_format == "spreadsheet"
    assert doc.parser_backend == "csv-reader"
    assert "QTE-2026-01521" in doc.full_text
    assert "12000" in doc.full_text
    assert "|" in doc.full_text  # rendered as a markdown pipe table
    # one page with a real table for the L2 line-item extractor
    assert doc.pages and doc.pages[0].tables


def test_router_dispatches_csv(tmp_path):
    f = tmp_path / "x.csv"
    f.write_text("a,b\n1,2\n")
    doc = route_parse(f)
    assert doc.file_format == "spreadsheet" and "1" in doc.full_text


def _write_workbook(path, cells):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for coord, value in cells.items():
        ws[coord] = value
    wb.save(path)


def test_xlsx_uncached_sum_formula_is_evaluated(tmp_path):
    # Programmatically-generated workbooks carry formulas with NO cached
    # values; the totals must still appear in full_text.
    f = tmp_path / "q.xlsx"
    _write_workbook(f, {
        "A1": "Line item", "B1": "Amount",
        "A2": "Platform licence", "B2": 100,
        "A3": "Support", "B3": 200,
        "A4": "Total", "B4": "=SUM(B2:B3)",
    })
    doc = parse_xlsx(f)
    assert "300" in doc.full_text


def test_xlsx_chained_arithmetic_formulas_are_evaluated(tmp_path):
    # =E2*0.2 referencing a cell that is itself a formula (VAT-style chain).
    f = tmp_path / "q.xlsx"
    _write_workbook(f, {
        "B2": 1000, "B3": 326,
        "E2": "=SUM(B2:B3)",   # 1326
        "E3": "=E2*0.2",       # 265.2
        "E4": "=E2*1.2",       # 1591.2
    })
    doc = parse_xlsx(f)
    assert "1326" in doc.full_text
    assert "265.2" in doc.full_text
    assert "1591.2" in doc.full_text


def test_xlsx_unsupported_formula_stays_blank(tmp_path):
    # Anything we cannot safely evaluate must render as blank — never the
    # formula source text, and never a fabricated number.
    f = tmp_path / "q.xlsx"
    _write_workbook(f, {"A1": "x", "B1": '=VLOOKUP(A1,Sheet2!A:B,2,0)'})
    doc = parse_xlsx(f)
    assert "VLOOKUP" not in doc.full_text


def test_xlsx_cached_values_still_win(tmp_path):
    # A workbook whose formulas DO carry cached results must parse exactly as
    # before (accuracy rule: evaluation only fills gaps, never overrides).
    import openpyxl
    f = tmp_path / "q.xlsx"
    _write_workbook(f, {"B2": 100, "B3": 200, "B4": "=SUM(B2:B3)"})
    # Simulate Excel having saved a cached result that disagrees with our math
    wb = openpyxl.load_workbook(f)
    wb.active["B4"] = 999  # plain value, as a cached-result stand-in
    wb.save(f)
    doc = parse_xlsx(f)
    assert "999" in doc.full_text


DEMO_XLSX = ("UI Improvements/ProcureIQ_Demo_Pack/quotes/06_SaaS/"
             "Aureus_Workflow_Ltd_V1.xlsx")


@pytest.mark.skipif(not os.path.exists(DEMO_XLSX), reason="demo pack not present")
def test_demo_pack_quote_totals_become_readable():
    # The exact file class that failed in session ses-20260730-UJF3:
    # Year-1 total ex-VAT = SUM(B16:B19) = 1326000, incl VAT = 1591200.
    doc = parse_xlsx(DEMO_XLSX)
    assert "1326000" in doc.full_text
    assert "1591200" in doc.full_text


@pytest.mark.skipif(not os.path.exists(XLSX), reason="sample xlsx not present")
def test_xlsx_parses_real_quote():
    doc = parse_xlsx(XLSX)
    assert doc.file_format == "spreadsheet"
    assert doc.parser_backend == "openpyxl"
    assert len(doc.full_text) > 50
    assert doc.pages  # at least one sheet


@pytest.mark.skipif(not os.path.exists(XLSX), reason="sample xlsx not present")
def test_router_no_longer_raises_on_xlsx():
    # Previously raised ValueError("unsupported file format: .xlsx")
    doc = route_parse(XLSX)
    assert doc.file_format == "spreadsheet"
