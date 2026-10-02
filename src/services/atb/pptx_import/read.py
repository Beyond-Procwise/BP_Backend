"""One pass over a .pptx, into plain frozen dataclasses.

Everything downstream measures THIS, never python-pptx objects, for three reasons: the
measurement modules stay pure and trivially testable, a shape's geometry is converted to inches
exactly once, and groups are flattened here so no later module can forget to walk into them.

Theme values are READ AND RECORDED, never used. On the reference deck the theme is stock Office
(Calibri Light, #4472C4) and has nothing to do with what the deck looks like — design §2.
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.util import Emu

_SRGB = re.compile(r'<a:srgbClr val="([0-9A-Fa-f]{6})"')
_LANG = re.compile(r'lang="([A-Za-z-]+)"')


class DeckUnreadable(Exception):
    """The file is not a usable presentation. Carries the reason, for the API to pass on."""


@dataclass(frozen=True)
class Box:
    x: float
    y: float
    w: float
    h: float


@dataclass(frozen=True)
class Run:
    text: str
    size_pt: float | None
    font: str | None
    colour: str | None
    bold: bool
    lang: str | None


@dataclass(frozen=True)
class Shape:
    kind: str
    box: Box
    runs: tuple[Run, ...]
    fill: str | None
    line: str | None
    table: tuple[tuple[str, ...], ...] | None
    slide: int
    name: str


@dataclass(frozen=True)
class Deck:
    width_in: float
    height_in: float
    slides: tuple[tuple[Shape, ...], ...]
    theme: dict[str, Any]
    chart_series_colours: tuple[str, ...]
    run_langs: dict[str, int]


def _inches(value) -> float:
    return 0.0 if value is None else round(Emu(int(value)).inches, 3)


def _hex(value: str | None) -> str | None:
    return None if value is None else '#' + value.upper()


def _fill_of(element) -> str | None:
    xml = element.xml
    start = xml.find('<a:solidFill>')
    if start < 0:
        return None
    match = _SRGB.search(xml, start)
    return _hex(match.group(1)) if match else None


def _line_of(element) -> str | None:
    xml = element.xml
    start = xml.find('<a:ln')
    if start < 0:
        return None
    match = _SRGB.search(xml, start)
    return _hex(match.group(1)) if match else None


def _kind_of(shape) -> str:
    if getattr(shape, 'has_table', False):
        return 'table'
    if getattr(shape, 'has_chart', False):
        return 'chart'
    name = str(getattr(shape, 'shape_type', '') or '')
    if 'PICTURE' in name:
        return 'picture'
    if 'PLACEHOLDER' in name:
        return 'placeholder'
    if shape.has_text_frame:
        return 'text'
    return 'auto'


def _runs_of(shape, langs: dict[str, int]) -> tuple[Run, ...]:
    if not shape.has_text_frame:
        return ()
    out: list[Run] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            lang = None
            match = _LANG.search(run._r.xml)
            if match:
                lang = match.group(1)
                langs[lang] = langs.get(lang, 0) + 1
            colour = None
            try:
                if run.font.color is not None and run.font.color.type is not None:
                    rgb = run.font.color.rgb
                    colour = _hex(str(rgb)) if rgb is not None else None
            except (AttributeError, TypeError, ValueError):
                colour = None
            out.append(Run(
                text=run.text or '',
                size_pt=float(run.font.size.pt) if run.font.size else None,
                font=run.font.name,
                colour=colour,
                bold=bool(run.font.bold),
                lang=lang,
            ))
    return tuple(out)


def _table_of(shape) -> tuple[tuple[str, ...], ...] | None:
    if not getattr(shape, 'has_table', False):
        return None
    return tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows)


def _walk(shapes, slide_no: int, langs: dict[str, int], out: list[Shape]) -> None:
    for shape in shapes:
        if 'GROUP' in str(getattr(shape, 'shape_type', '') or ''):
            # Review Focus 1: a group's children are invisible unless walked. Flattened here so
            # that no measurement module has to remember to do it.
            _walk(shape.shapes, slide_no, langs, out)
            continue
        out.append(Shape(
            kind=_kind_of(shape),
            box=Box(_inches(shape.left), _inches(shape.top),
                    _inches(shape.width), _inches(shape.height)),
            runs=_runs_of(shape, langs),
            fill=_fill_of(shape._element),
            line=_line_of(shape._element),
            table=_table_of(shape),
            slide=slide_no,
            name=str(shape.name or ''),
        ))


def _theme_of(data: bytes) -> dict[str, Any]:
    """Read from the zip rather than through python-pptx's relationship graph: the theme is only
    ever recorded, so the simplest reliable route is the right one."""
    with zipfile.ZipFile(BytesIO(data)) as archive:
        names = [n for n in archive.namelist() if n.startswith('ppt/theme/')]
        if not names:
            return {'colours': {}, 'fonts': {}}
        xml = archive.read(sorted(names)[0]).decode('utf-8', 'ignore')
    colours = dict(re.findall(
        r'<a:(dk1|dk2|lt1|lt2|accent[1-6])>.*?val="([0-9A-Fa-f]{6})"', xml))
    fonts = dict(re.findall(r'<a:(majorFont|minorFont)>\s*<a:latin typeface="([^"]*)"', xml))
    return {'colours': {k: _hex(v) for k, v in colours.items()}, 'fonts': fonts}


def _chart_series_colours(data: bytes) -> tuple[str, ...]:
    seen: list[str] = []
    with zipfile.ZipFile(BytesIO(data)) as archive:
        for name in sorted(n for n in archive.namelist() if n.startswith('ppt/charts/chart')):
            xml = archive.read(name).decode('utf-8', 'ignore')
            for match in _SRGB.finditer(xml):
                colour = _hex(match.group(1))
                if colour not in seen:
                    seen.append(colour)
    return tuple(seen)


def read_deck(data: bytes) -> Deck:
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise DeckUnreadable(f'the file could not be opened as a presentation: {exc}') from exc
    if not len(presentation.slides):
        raise DeckUnreadable('the presentation has no slides')
    langs: dict[str, int] = {}
    slides: list[tuple[Shape, ...]] = []
    for index, slide in enumerate(presentation.slides, 1):
        found: list[Shape] = []
        _walk(slide.shapes, index, langs, found)
        slides.append(tuple(found))
    try:
        theme = _theme_of(data)
    except Exception:
        theme = {'colours': {}, 'fonts': {}}
    try:
        series = _chart_series_colours(data)
    except Exception:
        series = ()
    return Deck(
        width_in=round(Emu(int(presentation.slide_width)).inches, 3),
        height_in=round(Emu(int(presentation.slide_height)).inches, 3),
        slides=tuple(slides),
        theme=theme,
        chart_series_colours=series,
        run_langs=langs,
    )
