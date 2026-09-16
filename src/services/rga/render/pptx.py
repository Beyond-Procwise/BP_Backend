"""The deck. Deterministic down to the byte.

WHY THE ZIP IS REWRITTEN

A .pptx is a zip, and ``python-pptx`` writes each entry with the wall clock in
its local header. Two renders of identical content therefore differ in the
bytes, which would make DoD12's "regenerate byte-identically" impossible to
satisfy for reasons that have nothing to do with the report. Every entry is
rewritten here with a fixed timestamp, in the original order. Measured: without
this, two renders a second apart differ; with it, they are identical.

The document properties are pinned for the same reason, and they are pinned to
values that describe the pack rather than the clock — so the file's own metadata
says which snapshot it came from.

THE BADGE RULE, AND WHY IT IS A LOCALITY RULE

Every drawing of a fact carries its footnote marker AND its badge **in the same
text frame**. Not somewhere on the slide, and not somewhere in the deck — in the
same frame.

This was a bug first. Badges were drawn only by ``_metric``, so a figure that
appeared solely in a table rendered as a bare "—" with nothing saying it was
unassessed rather than zero; and the post-check, which searched the whole deck
for the badge string, passed it because a *different* fact's badge had put the
word on the page. A guard satisfied by another figure's disclosure is not a
guard. The invariant below is what makes it checkable, and
``postcheck`` enforces it per fact rather than per deck.

WHAT THE RENDERER DECIDES: layout, and nothing else. Every figure printed comes
from ``FactEntry.display`` — the shared formatter — so the post-check's token
comparison is against the same spelling the reader sees. Badges and footnotes
are drawn here, unconditionally when the Style Brief asks for them, and the
origin badge is drawn whether it asks or not.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime
from typing import List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from src.services.analytics.models import Confidence
from src.services.rga.models import (
    ChartBlock,
    FactEntry,
    FactPack,
    FindingListBlock,
    MetricBlock,
    NarrativeBlock,
    PLACEHOLDER,
    ReportAST,
    Section,
    TableBlock,
    canonical_hash,
)
from src.services.rga.render import RenderedArtefact
from src.services.rga.style import StyleBrief

RENDERER = "pptx"
# Bump when the drawing changes. A rebuild that produced different bytes under
# the same version would be a silent break of the reproducibility promise.
RENDERER_VERSION = "1.2.0"

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# The DOS epoch. Any fixed value works; this one is conventional for
# reproducible archives.
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FIXED_DOC_TIME = datetime(2000, 1, 1, 0, 0, 0)

_BLANK_LAYOUT = 6

_LEGACY = "LEGACY_UNVERIFIED"


def badge_text(entry: FactEntry, show_confidence: bool) -> str:
    """The disclosure a fact must carry wherever it is drawn.

    The origin badge does not depend on the style key. A LEGACY_UNVERIFIED
    figure that looks like a verified one is the specific misreading it exists
    to prevent, so switching confidence badges off must not switch it off too.
    """
    parts: List[str] = []
    if show_confidence:
        parts.append(entry.confidence.value)
    if entry.origin.value == _LEGACY:
        parts.append(_LEGACY)
    return " · ".join(parts)


def _hex(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def _textbox(slide, left, top, width, height, text, *, size, bold=False,
             colour="#0A2A43", font="DM Sans", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = _hex(colour)
    return box


def _substitute(text: str, pack: FactPack) -> str:
    """Replace ``{{F0042}}`` with the fact's rendered form.

    An unresolvable reference is left as the literal placeholder rather than
    blanked: a sentence with a visible ``{{F0099}}`` in it is obviously broken,
    where a silently emptied one reads as finished prose that lost a number.
    The post-check catches it either way.
    """
    def _one(match) -> str:
        entry = pack.fact(match.group(1))
        return entry.display if entry else match.group(0)

    return PLACEHOLDER.sub(_one, text)


def render(
    ast: ReportAST,
    pack: FactPack,
    brief: StyleBrief,
    *,
    title: str = "Executive procurement summary",
) -> RenderedArtefact:
    """Draw the deck and return it with its reproducibility record."""

    palette = brief.get("report.style.palette")
    body_font = brief.get("report.style.font.body")
    mono_font = brief.get("report.style.font.mono")
    show_badges = brief.get("report.style.show_confidence_badges")
    show_footnotes = brief.get("report.style.show_provenance_footnotes")

    ast_hash = canonical_hash(ast.model_dump(mode="json"))
    style_version = brief.version()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    coverage = _coverage_line(pack)
    _title_slide(prs, pack, brief, palette, body_font, mono_font, title,
                 ast_hash, style_version, coverage)

    for section in ast.sections:
        _section_slide(prs, section, pack, palette, body_font, mono_font,
                       show_badges, show_footnotes)

    _pin_properties(prs, pack, title)

    buffer = io.BytesIO()
    prs.save(buffer)
    content = _normalise_zip(buffer.getvalue())

    return RenderedArtefact(
        content=content,
        media_type=MEDIA_TYPE,
        renderer=RENDERER,
        renderer_version=RENDERER_VERSION,
        pack_id=pack.pack_id,
        pack_hash=pack.hash,
        style_version=style_version,
        ast_hash=ast_hash,
        style_provenance=dict(brief.provenance),
        style_disclosure=brief.disclosure(),
        coverage_disclosure=coverage,
    )


def _coverage_line(pack: FactPack) -> str:
    """How much of itself this report could not measure.

    Stated on the front page rather than discovered per-figure. RB6 arrived at
    the same rule for the same reason (ruling R10): a rolled-up count of what is
    unmeasured has to be visible, or the reader learns it one caption at a time
    after they have already believed the deck.
    """
    unmeasured = sum(1 for f in pack.facts if f.confidence is Confidence.UNASSESSED)
    if not unmeasured:
        return f"all {len(pack.facts)} measures assessed"
    return f"{unmeasured} of {len(pack.facts)} measures not assessed"


# --------------------------------------------------------------------------
# Slides
# --------------------------------------------------------------------------


def _title_slide(prs, pack, brief, palette, body_font, mono_font, title,
                 ast_hash, style_version, coverage) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])

    _band(slide, prs, palette["ocean_dark"], Inches(2.6))
    _textbox(slide, Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.9), title,
             size=34, bold=True, colour="#FFFFFF", font=body_font)

    scope = pack.scope
    scope_line = " · ".join(str(v) for v in (
        scope.get("period_label"), scope.get("currency"),
        f"as at {pack.as_of}",
    ) if v)
    _textbox(slide, Inches(0.8), Inches(1.8), Inches(11.7), Inches(0.5), scope_line,
             size=14, colour="#FFFFFF", font=mono_font)

    incomplete = "not assessed" in coverage
    _textbox(slide, Inches(0.8), Inches(3.1), Inches(11.7), Inches(0.4), coverage,
             size=16, bold=True,
             colour=palette["warm"] if incomplete else palette["ok"], font=body_font)

    _textbox(slide, Inches(0.8), Inches(3.7), Inches(11.7), Inches(0.4),
             brief.disclosure(), size=10, colour=palette["mute"], font=body_font)

    stamp = (f"pack {pack.pack_id} · hash {pack.hash[:12]} · "
             f"style {style_version} · ast {ast_hash[:12]} · "
             f"renderer {RENDERER}/{RENDERER_VERSION}")
    _textbox(slide, Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.4), stamp,
             size=9, colour=palette["mute"], font=mono_font)


def _band(slide, prs, colour: str, height) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), prs.slide_width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(colour)
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.text_frame.text = ""


def _section_slide(prs, section: Section, pack, palette, body_font, mono_font,
                   show_badges, show_footnotes) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])

    _textbox(slide, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.6),
             section.title, size=26, bold=True, colour=palette["ocean_dark"],
             font=body_font)

    top = Inches(1.5)
    footnotes: List[str] = []

    for block in section.blocks:
        if isinstance(block, MetricBlock):
            top = _metric(slide, block, pack, palette, body_font, mono_font,
                          top, show_badges, footnotes)
        elif isinstance(block, NarrativeBlock):
            top = _narrative(slide, block, pack, palette, body_font, top,
                             show_badges, footnotes)
        elif isinstance(block, TableBlock):
            top = _table(slide, block, pack, palette, body_font, mono_font, top,
                         show_badges, footnotes)
        elif isinstance(block, ChartBlock):
            top = _chart_as_bars(slide, block, pack, palette, body_font,
                                 mono_font, top, show_badges, footnotes)
        elif isinstance(block, FindingListBlock):
            top = _findings(slide, block, pack, palette, body_font, top)

    if show_footnotes and footnotes:
        _textbox(slide, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.9),
                 "  ".join(footnotes), size=8, colour=palette["mute"],
                 font=body_font)


def _footnote_for(entry: FactEntry, show_badges: bool, footnotes: List[str]) -> str:
    """One numbered footnote per fact, reused if the fact appears twice.

    The badge is carried here as well as at the point of use, so the footnote
    line alone answers "how well is this figure known".
    """
    marker = f"[{entry.fact_id}]"
    badge = badge_text(entry, show_badges)
    parts = [marker] + ([badge] if badge else []) + [
        entry.label, entry.derivation, entry.provenance_id]
    note = " · ".join(parts)
    if note not in footnotes:
        footnotes.append(note)
    return marker


def _stamp(entry: FactEntry, show_badges: bool) -> str:
    """Marker and badge, together. The unit the post-check looks for."""
    badge = badge_text(entry, show_badges)
    return f"[{entry.fact_id}] · {badge}" if badge else f"[{entry.fact_id}]"


def _tone(entry: FactEntry, palette) -> str:
    if entry.confidence is Confidence.UNASSESSED:
        return palette["warm"]
    if entry.origin.value == _LEGACY:
        return palette["warm"]
    return palette["ok"]


def _metric(slide, block: MetricBlock, pack, palette, body_font, mono_font,
            top, show_badges, footnotes):
    entry = pack.fact(block.fact_ref)
    if entry is None:
        _textbox(slide, Inches(0.8), top, Inches(11.7), Inches(0.4),
                 f"[unresolved fact {block.fact_ref}]", size=14,
                 colour=palette["warm"], font=body_font)
        return top + Inches(0.5)

    primary = block.emphasis == "primary"
    _textbox(slide, Inches(0.8), top, Inches(6.5), Inches(0.35), entry.label,
             size=12, colour=palette["mute"], font=body_font)
    _textbox(slide, Inches(0.8), top + Inches(0.32), Inches(6.5), Inches(0.6),
             entry.display, size=30 if primary else 22, bold=True,
             colour=palette["ink"], font=mono_font)

    # The caption is the disclosure: reference and badge in one frame, directly
    # under the figure they qualify.
    _footnote_for(entry, show_badges, footnotes)
    _textbox(slide, Inches(0.8), top + Inches(0.92), Inches(6.5), Inches(0.3),
             _stamp(entry, show_badges), size=10, bold=True,
             colour=_tone(entry, palette), font=body_font)

    return top + Inches(1.35)


def _narrative(slide, block: NarrativeBlock, pack, palette, body_font, top,
               show_badges, footnotes):
    """Prose, and a stamp for every fact it rests on — quoted or not.

    A sentence may rely on a measured fact without printing its value: "No
    opportunities were identified" rests on a count of zero without quoting it,
    and that is better writing than forcing the figure into the line. But the
    reliance still has to be visible, or the claim is unsourced on the page.

    So the stamps come from the declared ``fact_refs`` as well as the inline
    placeholders. Drawing only the placeholders left facts referenced by the
    report and drawn nowhere in it — which the post-check duly caught, and was
    right to.
    """
    text = _substitute(block.text, pack)
    stamps: List[str] = []
    # Declared order first, then anything quoted but not declared, so the line
    # is stable across runs.
    seen = set()
    for ref in list(block.fact_refs) + block.placeholders():
        if ref in seen:
            continue
        seen.add(ref)
        entry = pack.fact(ref)
        if entry is not None:
            _footnote_for(entry, show_badges, footnotes)
            stamps.append(_stamp(entry, show_badges))

    _textbox(slide, Inches(0.8), top, Inches(11.7), Inches(0.6), text, size=14,
             colour=palette["ink"], font=body_font)
    if stamps:
        # Prose stays clean; the disclosure sits on its own line beneath it,
        # still one frame per fact-and-badge pair.
        _textbox(slide, Inches(0.8), top + Inches(0.5), Inches(11.7), Inches(0.3),
                 "   ".join(stamps), size=9, colour=palette["mute"],
                 font=body_font)
    return top + Inches(1.0 if stamps else 0.7)


def _table(slide, block: TableBlock, pack, palette, body_font, mono_font, top,
           show_badges, footnotes):
    rows = len(block.rows) + 1
    cols = len(block.columns)
    width = Inches(11.7)
    height = Inches(0.35) * rows
    shape = slide.shapes.add_table(rows, cols, Inches(0.8), top, width, height)
    table = shape.table

    for c, label in enumerate(block.columns):
        cell = table.cell(0, c)
        cell.text = label
        _style_cell(cell, palette["ocean_dark"], "#FFFFFF", body_font, bold=True)

    for r, row in enumerate(block.rows, start=1):
        for c, raw in enumerate(row):
            entry = pack.fact(raw)
            if entry is not None:
                _footnote_for(entry, show_badges, footnotes)
                # The badge travels with the figure. Without this an unassessed
                # value in a table read as a bare dash, indistinguishable from
                # a measured nothing.
                text = f"{entry.display} {_stamp(entry, show_badges)}"
                font, colour = mono_font, _tone(entry, palette)
            else:
                text, font, colour = raw, body_font, palette["ink"]
            cell = table.cell(r, c)
            cell.text = text
            _style_cell(cell, palette["paper"], colour, font)

    return top + height + Inches(0.3)


def _style_cell(cell, fill: str, colour: str, font: str, *, bold=False) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _hex(fill)
    for para in cell.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(11)
            run.font.bold = bold
            run.font.name = font
            run.font.color.rgb = _hex(colour)


def _chart_as_bars(slide, block: ChartBlock, pack, palette, body_font, mono_font,
                   top, show_badges, footnotes):
    """A chart drawn as labelled bars.

    Deliberately not a native pptx chart: an embedded chart carries its own
    workbook part, and that part is where a figure could sit that the post-check
    — which reads text frames — would never see. Bars keep every number in the
    text layer, where it is checkable.
    """
    from pptx.enum.shapes import MSO_SHAPE

    entries = [
        (series.label, pack.fact(ref))
        for series in block.series for ref in series.fact_refs
    ]
    measured = [e for _, e in entries if e is not None and e.value is not None]
    peak = max((abs(e.value) for e in measured), default=None)

    y = top
    for label, entry in entries:
        if entry is None:
            continue
        _textbox(slide, Inches(0.8), y, Inches(3.2), Inches(0.3),
                 f"{label} — {entry.label}", size=11, colour=palette["mute"],
                 font=body_font)
        if entry.value is not None and peak:
            span = Inches(6.0) * float(abs(entry.value)) / float(peak)
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), y,
                                         max(int(span), Emu(1)), Inches(0.24))
            bar.fill.solid()
            bar.fill.fore_color.rgb = _hex(palette["ocean"])
            bar.line.fill.background()
            bar.shadow.inherit = False
            bar.text_frame.text = ""
        _footnote_for(entry, show_badges, footnotes)
        _textbox(slide, Inches(10.4), y, Inches(2.5), Inches(0.3),
                 f"{entry.display} {_stamp(entry, show_badges)}", size=10,
                 bold=True, colour=_tone(entry, palette), font=mono_font)
        y = y + Inches(0.42)

    return y + Inches(0.2)


def _findings(slide, block: FindingListBlock, pack, palette, body_font, top):
    wanted = set(block.finding_refs)
    listed = [f for f in pack.findings if not wanted or f.finding_id in wanted]
    if not listed:
        _textbox(slide, Inches(0.8), top, Inches(11.7), Inches(0.4),
                 "No findings raised against this pack.", size=13,
                 colour=palette["mute"], font=body_font)
        return top + Inches(0.5)

    y = top
    for found in listed:
        _textbox(slide, Inches(0.8), y, Inches(11.7), Inches(0.4),
                 f"{found.severity.value} · {found.code.value} — {found.detail}",
                 size=11,
                 colour=palette["warm"] if found.blocks_release else palette["mute"],
                 font=body_font)
        y = y + Inches(0.38)
    return y + Inches(0.2)


# --------------------------------------------------------------------------
# Determinism
# --------------------------------------------------------------------------


def _pin_properties(prs, pack: FactPack, title: str) -> None:
    """Metadata that describes the snapshot, not the clock."""
    core = prs.core_properties
    core.title = title
    core.author = pack.generated_by
    core.last_modified_by = pack.generated_by
    core.comments = f"Fact Pack {pack.pack_id} · {pack.hash}"
    core.revision = 1
    core.created = _FIXED_DOC_TIME
    core.modified = _FIXED_DOC_TIME


def _normalise_zip(raw: bytes) -> bytes:
    """Rewrite every entry with a fixed timestamp, preserving order.

    Without this the same deck rendered twice differs, because each entry's
    header carries the wall clock at the moment it was written.
    """
    source = io.BytesIO(raw)
    target = io.BytesIO()
    with zipfile.ZipFile(source) as zin, \
            zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            pinned = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            pinned.compress_type = zipfile.ZIP_DEFLATED
            pinned.external_attr = info.external_attr
            pinned.internal_attr = info.internal_attr
            pinned.create_system = 0
            zout.writestr(pinned, zin.read(info.filename))
    return target.getvalue()


def extract_text(content: bytes) -> List[str]:
    """Every string the deck actually shows. What the post-check reads.

    Tables are walked explicitly: a table cell's text lives on the cell's own
    frame, not on the graphic frame's, so a check that only walked
    ``has_text_frame`` would read past every figure in every table.
    """
    found: List[str] = []
    prs = Presentation(io.BytesIO(content))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                found.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            found.append(cell.text)
    return found
