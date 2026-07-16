# scripts/export_benchmark_fixtures.py
"""One-time exporter: Benchmark Calculations.xlsx -> tests/fixtures/benchmark/golden.json.

Run from the repo root:  python scripts/export_benchmark_fixtures.py

Captures every input (quote lines, benchmark points, lookup tables) and every
computed output column from the Pricing Calculation sheet, so the parity suite
is hermetic — openpyxl is needed only here, never at test time.
"""
import json
import warnings
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Benchmark Calculations" / "Benchmark Calculations.xlsx"
OUT = ROOT / "tests" / "fixtures" / "benchmark" / "golden.json"

# Pricing Calculation computed columns (0-based tuple index -> fixture key).
EXPECTED_COLS = {
    22: "n_internal", 23: "n_external", 24: "n_total", 27: "total_weight",
    28: "simple_benchmark", 29: "weighted_benchmark", 30: "median_excel_buggy",
    31: "selected_benchmark", 32: "ref_quantity", 33: "avg_spec_score",
    34: "target_loc_index", 35: "avg_loc_index", 36: "avg_sla_score",
    37: "current_index", 38: "avg_hist_index",
    39: "volume_adjustment", 40: "spec_adjustment", 41: "location_adjustment",
    42: "sla_adjustment", 43: "inflation_adjustment",
    44: "final_benchmark", 45: "quoted_total", 46: "benchmark_total",
    47: "unit_variance_gbp", 48: "unit_variance_pct", 49: "total_cost_gap",
    50: "combined_factor", 51: "confidence",
}


def read_points(sheet, source):
    points = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        points.append({
            "benchmark_point_id": row[0], "source": source,
            "item_name": row[1], "uom": row[5], "currency": row[6],
            "include": row[18],
            "raw_unit_price": row[12], "source_weight": row[13],
            "specification_score": row[4], "location_cost_index": row[8],
            "sla_score": row[10], "historical_quantity": row[11],
            "index_value_at_price_date": row[17],
        })
    return points


def read_lookup(sheet):
    table = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        # Filters the trailing "HOW TO ..." note rows (non-numeric index col).
        if row[0] and isinstance(row[2], (int, float)):
            table[row[0]] = float(row[2])
    return table


def main():
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    points = read_points(wb["Internal Benchmark Data"], "internal") + \
        read_points(wb["External DB Benchmark Extract"], "external")

    quotes, expected = [], {}
    for row in wb["Pricing Calculation"].iter_rows(min_row=5, values_only=True):
        if row[0] is None:
            continue
        quotes.append({
            "deal_id": row[0], "item_name": row[1], "supplier_name": row[2],
            "quote_ref": row[3], "category": row[4], "quantity": row[5],
            "uom": row[6], "currency": row[7], "location": row[8],
            "requested_spec_score": row[9], "service_level": row[10],
            "requested_sla_score": row[11], "index_id": row[12],
            "quoted_unit_price": row[13], "delivery_cost": row[14],
            "implementation_cost": row[15], "support_cost": row[16],
            "risk_premium": row[17], "discount_rebate": row[18],
            "supplier_risk_penalty": row[19], "contract_risk_penalty": row[20],
            "strategic_supplier_bonus": row[21],
        })
        expected[row[0]] = {key: row[i] for i, key in EXPECTED_COLS.items()}

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({
        "location_index_table": read_lookup(wb["Location Index Table"]),
        "index_table": read_lookup(wb["Index Table"]),
        "points": points, "quotes": quotes, "expected": expected,
    }, indent=2, default=str))
    print(f"wrote {OUT} ({len(points)} points, {len(quotes)} quotes)")


if __name__ == "__main__":
    main()
